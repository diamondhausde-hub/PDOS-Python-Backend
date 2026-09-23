import io
import csv
from fpdf import FPDF
from openpyxl import Workbook
from typing import List, Dict, Any

def build_rep_pdf(data: Dict[str, Any]) -> bytes:
    """
    Builds a plain, clean PDF layout for a rep report using fpdf2.
    """
    pdf = FPDF()
    pdf.add_page()
    
    import os
    logo_path = "static/report_logo.png"
    if os.path.exists(logo_path):
        pdf.image(logo_path, x=10, y=8, w=20)
    
    # Title
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Rep Performance Report", new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(10)
    
    # Summary KPIs
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Key Performance Indicators", new_x="LMARGIN", new_y="NEXT")
    
    pdf.set_font("Helvetica", "", 12)
    pdf.cell(0, 8, f"Total Revenue: ${data.get('revenue', 0):.2f}", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 8, f"Visits Completed: {data.get('visits_completed', 0)}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)
    
    # Targets
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Targets Progress", new_x="LMARGIN", new_y="NEXT")
    
    pdf.set_font("Helvetica", "", 12)
    targets = data.get("targets", [])
    if not targets:
        pdf.cell(0, 8, "No targets defined for this period.", new_x="LMARGIN", new_y="NEXT")
    else:
        for t in targets:
            # Assuming target has title, achieved_qty, target_qty
            name = t.get('title', 'Target')
            achieved = t.get('achieved_qty', 0)
            target_qty = t.get('target_qty', 1)
            target_qty = target_qty if target_qty > 0 else 1
            pdf.cell(0, 8, f"- {name}: {achieved} / {target_qty} ({achieved/target_qty*100:.1f}%)", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)
    
    # Top Products
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Top Selling Products", new_x="LMARGIN", new_y="NEXT")
    
    pdf.set_font("Helvetica", "", 12)
    top_products = data.get("top_products", [])
    if not top_products:
        pdf.cell(0, 8, "No product sales recorded in this period.", new_x="LMARGIN", new_y="NEXT")
    else:
        for p in top_products:
            # Assuming product has product_name, units_sold, total_revenue
            name = p.get('product_name', 'Product')
            units = p.get('units_sold', 0)
            rev = p.get('total_revenue', 0.0)
            pdf.cell(0, 8, f"- {name}: {units} units (${rev:.2f})", new_x="LMARGIN", new_y="NEXT")
            
    return bytes(pdf.output())

def build_excel(rows: List[Dict[str, Any]]) -> bytes:
    """
    Builds an Excel file from a list of dictionaries.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Export"
    
    import os
    from openpyxl.drawing.image import Image as OpenpyxlImage
    logo_path = "static/report_logo.png"
    if os.path.exists(logo_path):
        img = OpenpyxlImage(logo_path)
        img.height = 60
        img.width = 60
        ws.add_image(img, "A1")
        
        # Offset rows to make space for the logo
        ws.append([])
        ws.append([])
        ws.append([])
    
    if rows:
        # Headers
        headers = list(rows[0].keys())
        ws.append(headers)
        
        # Rows
        for row in rows:
            ws.append([row.get(h) for h in headers])
            
    # Save to bytes
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

def build_csv(rows: List[Dict[str, Any]]) -> bytes:
    """
    Builds a CSV file from a list of dictionaries.
    """
    out = io.StringIO()
    if rows:
        headers = list(rows[0].keys())
        writer = csv.DictWriter(out, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)
    return out.getvalue().encode('utf-8')
